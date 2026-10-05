from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'report' / 'AI_Travel_Agent_College_Report_Revised.docx'
OUT.parent.mkdir(exist_ok=True)
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = s.top_margin = Cm(2.5)
s.bottom_margin = Cm(2)
for name, size in [('Normal',12),('Title',16),('Heading 1',16),('Heading 2',14),('Caption',10)]:
    st=d.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().append(OxmlElement('w:rFonts')) if st.element.get_or_add_rPr().find(qn('w:rFonts')) is None else None
    rf=st.element.get_or_add_rPr().find(qn('w:rFonts'))
    for key in list(rf.attrib):
        if 'theme' in key.lower(): del rf.attrib[key]
    for a in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+a),'Times New Roman')
d.styles['Normal'].paragraph_format.line_spacing=1.3
d.styles['Normal'].paragraph_format.space_after=Pt(6)
for name in ['Heading 1','Heading 2']:
    d.styles[name].font.bold=True
    d.styles[name].paragraph_format.space_before=Pt(6)
    d.styles[name].paragraph_format.space_after=Pt(10)
    d.styles[name].paragraph_format.keep_with_next=True
d.styles['Caption'].font.italic=True
d.styles['Caption'].paragraph_format.line_spacing=1
for st in [d.styles['Title'],d.styles['Normal']]:
    pr=st.element.find(qn('w:pPr'))
    if pr is not None:
        for border in list(pr.findall(qn('w:pBdr'))): pr.remove(border)

def spacing(run):
    rp=run._element.get_or_add_rPr(); e=OxmlElement('w:spacing'); e.set(qn('w:val'),'10'); rp.append(e)

def p(text, center=False):
    a=d.add_paragraph(); a.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    a.add_run(text); return a

def h(text): d.add_paragraph(text,'Heading 2')
def page(): d.add_page_break()
def chapter(n,title):
    page(); a=d.add_paragraph(f'CHAPTER {n}\n{title.upper()}','Heading 1'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER

def table(rows, headers=('Module or File','Implementation responsibility'), widths=(5,11)):
    t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
    for i,x in enumerate(headers): t.rows[0].cells[i].text=x
    for row in rows:
        c=t.add_row().cells
        for i,x in enumerate(row): c[i].text=str(x)
    for j,row in enumerate(t.rows):
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]); c.vertical_alignment=1
            pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            if j==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7E6E6'); pr.append(sh)
            for a in c.paragraphs:
                a.paragraph_format.line_spacing=1.1; a.paragraph_format.space_after=Pt(4); a.paragraph_format.space_before=Pt(4)
                for r in a.runs: r.font.size=Pt(12); r.font.bold=j==0
        rp=row._tr.get_or_add_trPr(); e=OxmlElement('w:cantSplit'); rp.append(e)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def gap(label,height=13):
    a=p(label,True); a.paragraph_format.space_before=Pt(12)
    b=d.add_paragraph(); b.paragraph_format.space_after=Pt(0); b.paragraph_format.line_spacing=Pt(height*28.346)
    b.add_run(' ')

def caption(text):
    a=d.add_paragraph(text,'Caption'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER

# Native Word equation constructors. All displayed mathematics remains editable.
def mt(text):
    r=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=text; r.append(t); return r
def frac(a,b):
    e=OxmlElement('m:f'); n=OxmlElement('m:num'); de=OxmlElement('m:den')
    n.append(mt(a)); de.append(mt(b)); e.extend([n,de]); return e
def sub(a,b):
    e=OxmlElement('m:sSub'); x=OxmlElement('m:e'); y=OxmlElement('m:sub'); x.append(mt(a)); y.append(mt(b)); e.extend([x,y]); return e
def sup(a,b):
    e=OxmlElement('m:sSup'); x=OxmlElement('m:e'); y=OxmlElement('m:sup'); x.append(mt(a)); y.append(mt(b)); e.extend([x,y]); return e
def eq(*parts):
    a=d.add_paragraph(); a.alignment=WD_ALIGN_PARAGRAPH.CENTER; a.paragraph_format.line_spacing=1.3
    a.paragraph_format.space_before=Pt(4); a.paragraph_format.space_after=Pt(4)
    om=OxmlElement('m:oMathPara'); math=OxmlElement('m:oMath')
    for part in parts: math.append(mt(part) if isinstance(part,str) else part)
    om.append(math); a._p.append(om)

# Cover
a=d.add_paragraph('AI TRAVEL AGENT','Title'); a.alignment=WD_ALIGN_PARAGRAPH.CENTER
a.paragraph_format.keep_with_next=False
for pr in [a._p.get_or_add_pPr(), d.styles['Title'].element.find(qn('w:pPr'))]:
    if pr is not None:
        for border in list(pr.findall(qn('w:pBdr'))): pr.remove(border)
a.paragraph_format.space_before=Pt(110)
p('College Project Report',True)
p('Artificial Intelligence and Natural Language Processing',True)
p('Conversational travel planning with constraint optimisation and route ordering',True)
p('October 2026',True)

page(); d.add_paragraph('TABLE OF CONTENTS','Heading 1')
topics=['Introduction','Objectives','System Description and NLP Approach','AI and NLP Pipeline and Architecture','Mathematical Formulation and Algorithms','Tech Stack and Packages','Database and Knowledge Base Structure','Modules and Implementation','Output Screenshots','Conclusion and Future Enhancement']
chapter_pages=[3,4,5,7,8,12,14,15,20,45]
for i,title in enumerate(topics,1):
    a=p(f'Chapter {i}   {title}\t{chapter_pages[i-1]}'); a.alignment=WD_ALIGN_PARAGRAPH.LEFT
    a.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    a.paragraph_format.space_after=Pt(10)
a=p('References and source materials\t46'); a.alignment=WD_ALIGN_PARAGRAPH.LEFT
a.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)

chapter(1,topics[0]); h('1.1 Background')
p('Planning a trip requires coordinating destinations, travel dates, transport, accommodation, food, sightseeing and expenses. These decisions depend on one another. A destination that appears attractive may be too far from the origin, an outdoor activity may be unsuitable during rain, and a day containing several long treks may be physically unrealistic. Searching separate websites does not automatically produce a consistent day-wise plan.')
h('1.2 Problem Statement')
p('Travellers usually express requirements informally, such as “Plan a relaxed three-day trip from Chennai to Ooty for two people.” A useful planning system must identify these details, ask for missing information, retrieve relevant places and turn preferences into a feasible schedule. It must also distinguish verified attraction facts from estimated duration, travel distance, cost and weather suitability.')
h('1.3 Project Overview')
p('The AI Travel Agent provides a conversational interface for personalised trip planning. Its React application communicates with a FastAPI backend and a Python planning agent. The agent combines natural language understanding, destination discovery, external location and weather services, activity classification, constraint-based scheduling and route optimisation. Trip state and conversation history are stored in SQLite so users can reopen saved sessions.')
h('1.4 Scope')
p('The project focuses on itinerary assistance, with curated Indian destinations and attraction data supplemented by OpenStreetMap. It supports destination selection, recommendations, day-wise plans, map views and budget estimates. OR-Tools assigns attractions to days, while route-ordering algorithms reduce travel within each day. Prices are planning assumptions rather than live booking quotations. Weather can be a short-range forecast or a historical climate estimate, depending on the trip dates. The resulting itinerary is a decision aid that travellers can review and edit before travel.')

chapter(2,topics[1]); h('2.1 Primary Objective')
p('The primary objective is to convert a traveller’s natural-language requirements into a personalised and explainable itinerary by combining NLP, travel data retrieval and optimisation.')
h('2.2 Specific Objectives')
objectives=[
('Understand travel requests','Extract origin, destination, dates, budget, traveller count, transport mode and preferences from conversational messages.'),
('Identify user intent','Recognise requests for recommendations, itinerary generation, weather, transport, budget information and changes to the plan.'),
('Discover relevant destinations','Use category and proximity search to suggest destinations when the user has not selected a specific place.'),
('Build feasible schedules','Assign attractions to days while respecting stop limits, daily effort, intensity, trek limits and supported opening-hour information.'),
('Improve route order','Use road-distance data when available and apply A* or nearest-neighbour with 2-opt to sequence daily visits.'),
('Personalise recommendations','Consider interests, pace, accessibility requirements, weather suitability and crowd-avoidance preferences.'),
('Present useful outputs','Display recommendations, daily timelines, stay options, budget estimates and geographic locations in the web interface.'),
('Preserve sessions and provenance','Save trip state and messages, and label attraction facts and estimated values with their sources.')]
for i,(title,body) in enumerate(objectives,1):
    a=p(f'{i}. {title}. {body}'); a.paragraph_format.space_after=Pt(7)

chapter(3,topics[2]); h('3.1 System Description')
p('The application has a browser interface, an API layer and a planning layer. React manages the conversation and trip views. FastAPI exposes session creation, chat, recommendations, itinerary movement and map endpoints. The Python agent owns the conversational state and invokes the tools needed to complete a travel request. Pydantic models provide a consistent representation of trip dates, destinations, preferences and generated outputs.')
h('3.2 Named Entity Recognition and Slot Extraction')
p('The NER module extracts structured trip slots before the agent decides its next action. It uses spaCy’s English model when available, with regular-expression rules that continue to work when the model is missing. The rules cover Indian date formats, amounts such as “40k INR”, durations, traveller counts, solo travel, origins and transport keywords. Place names are resolved through geocoding and destination-specific logic.')
p('For example, “From Chennai to Ooty for two people, budget 30000 INR, from 10/10/2026 to 12/10/2026” supplies an origin, destination, group size, budget and date interval. The extracted slots are merged into the session state. Missing or uncertain details are handled through further conversation rather than silently inventing values.')
h('3.3 Intent Classification')
p('Intent recognition has two layers. High-precision rules identify clear messages such as greetings, numbered destination choices and explicit scheduling commands. Other messages pass through a local classifier trained on labelled examples. Word unigrams and bigrams are combined with character n-grams using TF-IDF, and Logistic Regression predicts the intent and its confidence.')
p('Intent categories include providing trip details, requesting recommendations, building an itinerary, changing pace, removing a stop, asking about budget, confirming a destination, asking about transport, asking about weather, greeting and other requests. Rule matches are trusted directly; classifier predictions require sufficient confidence before they trigger a shortcut.')
page(); h('3.4 Semantic Routing and Tool Dispatch')
p('A hybrid router determines whether a message should use the deterministic NLP tool engine or the LLM conversation engine. It compares local FastEmbed sentence embeddings against example requests using cosine similarity and also considers intent confidence, message complexity and the current trip state. A TF-IDF fallback is available when embedding initialisation fails.')
p('The deterministic tool router checks whether the required state is available before calling a tool. A weather request, for example, needs a resolved destination and relevant dates. Simple commands can therefore be processed without an LLM tool-calling round. Open-ended comparisons, ambiguous requests and conversational repair are passed to the orchestrator’s LLM path.')
h('3.5 LLM Orchestration')
p('The orchestrator maintains conversation history, fits messages within a token budget and coordinates tool results with the selected language model. The shared client supports Groq, NVIDIA through an OpenAI-compatible API, and a local OpenAI-compatible service such as Ollama. Model selection and credentials are environment settings; the application does not require one fixed model for every deployment.')
p('Activity classification is another LLM-assisted task. It estimates visit duration, activity class, intensity and related attributes when stronger place-specific information is unavailable. Returned values are validated and bounded before scheduling. Cache reuse reduces repeated requests, while conservative defaults allow the planner to continue when classification fails.')
h('3.6 Data Grounding and Limitations')
p('Destination and attraction information comes from curated records and external services. The model does not serve as the sole source of attraction facts. Approved profiles, OpenStreetMap tags, category rules and model estimates are distinguished through provenance fields. Missing prices and accessibility remain unknown unless a source supplies them. NLP extraction can still misinterpret uncommon place names, indirect phrasing or unsupported language, so the interface and conversation should allow the traveller to inspect and correct the plan.')

chapter(4,topics[3])

chapter(5,topics[4]); h('5.1 Text Representation and Intent Prediction')
p('The classifier represents each user message as a vector of word and character features. For a term t in message q, the following expressions summarise the smoothed TF-IDF weighting used by the vectoriser. N is the number of training messages and df(t) is the number containing the term. The resulting feature vector is normalised before classification.')
eq(sub('idf','t'),' = ln(',frac('1 + N','1 + df(t)'),') + 1')
eq(sub('v','t,q'),' = ',sub('tf','t,q'),' × ',sub('idf','t'))
eq('P(y = k | v) = ',frac('exp(wₖᵀv + bₖ)','Σⱼ exp(wⱼᵀv + bⱼ)'))
p('The predicted intent is the class with the highest probability. The NLU entry point treats rule matches as high confidence and requires a classifier confidence of at least 0.70 for a high-confidence result. This threshold controls dispatch; it is not a measured accuracy score.')
h('5.2 Semantic Similarity')
eq('cos(u, v) = ',frac('u · v','‖u‖ ‖v‖'))
p('The semantic router compares a message embedding u with route exemplar embeddings v. Larger cosine similarity indicates stronger alignment with an exemplar. Similarity is combined with confidence and dialogue-state checks to choose the NLP or LLM path. The router uses pretrained local embeddings rather than training a language model from the trip database.')
h('5.3 Geographic Distance')
eq('a = sin²(Δφ / 2) + cos(φ₁) cos(φ₂) sin²(Δλ / 2)')
eq('d = 2R arcsin(√a)')
p('Here φ denotes latitude, λ denotes longitude, angles are in radians, and R is approximately 6371 km. Haversine distance supports nearby destination search and radius filtering. It measures straight-line surface distance; road-distance services are preferred for travel planning. The scheduler uses a labelled approximation of 1.4 times Haversine distance when its road matrix is unavailable.')

page(); h('5.4 Attraction Assignment Variables')
p('Let I be the set of candidate attractions and D the set of trip days. A binary variable xᵢ,ᵈ equals 1 when attraction i is assigned to day d and 0 otherwise. Let eᵢ be visit effort in minutes, qᵢ be intensity and tᵢ indicate a trek. Each attraction can be omitted if the complete pool cannot satisfy the traveller’s limits.')
eq(sub('x','i,d'),' ∈ {0, 1}')
eq(sub('∑','d ∈ D'),' ',sub('x','i,d'),' ≤ 1   for every attraction i')
eq(sub('∑','i ∈ I'),' ',sub('x','i,d'),' ≤ M   for every day d')
p('The first inequality prevents duplicate visits. The second limits the daily number of stops to M. Daily visit effort and intensity are bounded separately so a short sightseeing stop and a long trek are not treated as equivalent.')
eq(sub('∑','i ∈ I'),' ',sub('e','i'),' ',sub('x','i,d'),' ≤ ',sub('E','max'))
eq(sub('∑','i ∈ I'),' ',sub('q','i'),' ',sub('x','i,d'),' ≤ ',sub('Q','max'))
p('The implemented default effort caps are 240, 330 and 420 minutes for relaxed, balanced and packed pace. The documented intensity defaults are 5, 8 and 11 respectively. These limits constrain the planning model; they are not medical assessments of an individual traveller.')
h('5.5 Trek and Availability Constraints')
eq(sub('∑','i ∈ I'),' ',sub('t','i'),' ',sub('x','i,d'),' ≤ 1')
eq(sub('∑','d ∈ D'),' ',sub('∑','i ∈ I'),' ',sub('t','i'),' ',sub('x','i,d'),' ≤ ',sub('T','max'))
p('At most one trek is allowed per day, and the total number of treks is limited by the selected adventure level. A supported opening-hours expression that rules out a visit on a particular date forces xᵢ,ᵈ to zero. Extremely distant attraction pairs cannot share a day. Complex opening-hours expressions remain unverified rather than being interpreted as guaranteed availability.')

page(); h('5.6 CP SAT Objective Function')
p('OR-Tools CP-SAT chooses a feasible assignment using a weighted objective. The following notation summarises the implemented components. S is the total number of scheduled attractions, R is rain exposure, C is the condition penalty, P is pairwise geographic dispersion, H is nearby-place cohesion, L is stop-count imbalance, E is effort imbalance, V is activity diversity and A penalises excessive repetition of one activity class.')
eq('min J = −10000S + wᵣR + C + P − wₕH')
eq('                 + wₗL + wₑE − wᵥV + wₐA')
p('The large negative coefficient rewards scheduling attractions that satisfy the hard constraints. Rain and condition penalties discourage unsuitable day assignments. Cohesion favours grouping nearby places; dispersion discourages excessive separation. Load and effort terms distribute the workload across days. Diversity rewards a mixture of activity classes instead of repeatedly scheduling the same type of stop.')
eq('S = ',sub('∑','i ∈ I'),' ',sub('∑','d ∈ D'),' ',sub('x','i,d'))
eq('L = ',sub('∑','d ∈ D'),' |',sub('∑','i ∈ I'),' ',sub('x','i,d'),' − target|')
p('Pairwise co-assignment uses an auxiliary Boolean variable bᵢ,ⱼ,ᵈ. It is true only when attractions i and j are both scheduled on day d. The equivalent linear inequalities below let CP-SAT express pairwise rewards and penalties.')
eq('bᵢ,ⱼ,ᵈ ≤ xᵢ,ᵈ     and     bᵢ,ⱼ,ᵈ ≤ xⱼ,ᵈ')
eq('bᵢ,ⱼ,ᵈ ≥ xᵢ,ᵈ + xⱼ,ᵈ − 1')
h('5.7 Solver and Timeline Construction')
p('The solver runs within a time limit and accepts an optimal or feasible solution. A greedy assignment is used if CP-SAT cannot provide either. After assignment, geographic day ordering and within-day route sequencing produce arrival and departure times. Meals and travel durations are inserted into the timeline, and the result is checked for daily overruns. Current assignment operates on individual attractions with pairwise cohesion; geographic DBSCAN pre-clustering has been removed from this stage.')
p('Budget estimation is a separate calculation. The current CP-SAT assignment should not be described as a proof that the total trip cost is below the user’s budget. The displayed estimate helps the traveller compare the plan with available funds.')

page(); h('5.8 A Star Route Ordering')
p('For a day with at most eight stops, A* searches visit sequences that begin and end at the hotel node. A search state contains a visited-stop bitmask and the current stop. The priority combines the distance already travelled with a lower-bound estimate of the remaining route.')
eq('f(n) = g(n) + h(n)')
eq('h = min distance to U + MST(U) + min return distance')
p('U is the unvisited set. The heuristic combines a connection to U, a minimum spanning tree over U, and a connection back to the hotel. For the tree, the smaller of the two directed distances is used. Road-matrix distances determine route quality; an estimated matrix means the resulting order is optimal for that model rather than a guarantee of the shortest real-world road journey.')
h('5.9 Nearest Neighbour and Two Opt')
p('For larger daily groups, nearest-neighbour selects the next unvisited stop with the lowest travel cost. A 2-opt improvement step reverses route segments when doing so reduces the route cost. These heuristics reduce search effort but do not guarantee a globally optimal route. Manual itinerary movement preserves the user’s chosen order and retimes the affected day.')
h('5.10 Budget Estimation')
eq(sub('C','total'),' = ',sub('C','stay'),' + ',sub('C','food'),' + ',sub('C','local'),' + ',sub('C','entry'),' + ',sub('C','intercity'))
eq(sub('C','stay'),' = max(D − 1, 0) × 3500 × ⌈P / 2⌉')
eq(sub('C','food'),' = D × P × 2 × 400')
eq(sub('C','local'),' = D × 1500')
p('D denotes trip days and P denotes travellers. The implemented assumptions are INR 3500 per room-night, INR 400 per meal and INR 1500 per day for local transport. Known entry fees are combined with an INR 100 per-person allowance for each unknown-fee stop. Intercity transport adds a model-based fare or fuel estimate when the selected mode and distance are available. These are explicit assumptions, not live market prices.')

chapter(6,topics[5]); h('6.1 Application and NLP Technologies')
table([
('React and React DOM','Browser interface, conversation state and trip views.'),('Vite','Frontend development server and build tooling.'),('React Router','Navigation between landing and trip-builder pages.'),('Leaflet and React Leaflet','Interactive map components and geographic markers.'),('MapLibre GL and React Map GL','Vector map rendering and map integration.'),('Python','Planning logic, NLP modules and service clients.'),('FastAPI and Uvicorn','HTTP API endpoints and ASGI application server.'),('Pydantic','Trip-state schemas and request/response validation.'),('spaCy','Optional English named-entity recognition.'),('scikit-learn','TF-IDF features, Logistic Regression and router fallback.'),('FastEmbed','Local ONNX sentence embeddings for semantic routing.'),('Joblib','Persistence of the intent classifier.'),('Groq SDK','Configured Groq language-model access.'),('OpenAI SDK','OpenAI-compatible NVIDIA and local model access.')],('Technology or Package','Role in the project'))
p('The provider selected at runtime depends on environment configuration. The presence of a package in requirements does not establish that it is used by the active React and FastAPI application.')
page(); h('6.2 Optimisation Data and Supporting Packages')
table([
('OR-Tools','CP-SAT assignment of attractions to trip days.'),('SQLite and sqlite3','Local session state, chat history and trace storage.'),('JSON','Approved place profiles, external-response caches and structured state.'),('Requests','HTTP calls to geocoding, weather and travel-data services.'),('python-dotenv','Loading configured environment settings.'),('Scrapy','Optional official-place refresh with a review queue.'),('OpenStreetMap and Overpass','Geographic objects, attractions, stays, food and source tags.'),('Nominatim','Place-name geocoding and fallback location resolution.'),('GeoNames','Supplementary destination-discovery source.'),('Open-Meteo','Weather forecasts and historical climate estimates.'),('Ola Maps','Configured geocoding, road matrices, routes and elevation.'),('Tailwind CSS and PostCSS','Existing frontend styling and CSS processing.'),('ESLint','Frontend static analysis.'),('Streamlit Folium and streamlit-folium','Alternative Streamlit interface and map rendering.'),('google-genai','Declared dependency; shared active client supports other providers.')],('Technology or Service','Purpose or usage status'))
p('Road and elevation services require configuration. Cached responses and labelled estimates provide alternatives when live data is unavailable. The React frontend and FastAPI backend are the main application described in this report; Streamlit is an additional entry point in the repository.')

chapter(7,topics[6])

chapter(8,topics[7]); h('8.1 Project Folder Structure')
gap('[Insert screenshot of the project folder structure here]',10)
caption('Figure 8.1  Project folder structure')
p('The folder structure separates the agent, backend server, web interface, data sources and tests. The agent directory contains NLP and planning modules. The server directory exposes API endpoints, while web contains the React application. The data, scripts and test directories support knowledge maintenance and validation.')
h('8.2 Module Organisation')
p('The following tables list the application’s source modules individually, grouped by responsibility. Supporting data, startup scripts and test modules are listed separately. Generated caches, dependency folders and bytecode are runtime artefacts rather than implementation modules.')
p('Windows launch helpers include start.bat, start.ps1, stop.bat, run_backend.bat, run_frontend.bat, free_ports.ps1 and wait_for_backend.ps1. Requirements and package manifests declare Python and web dependencies. Test files document intended checks; their presence alone does not establish that all scenarios pass.')
page(); h('8.3 Conversation NLP and State Modules')
table([
('agent/__init__.py','Initialises the Python agent package.'),('agent/orchestrator.py','Maintains dialogue history and coordinates routing, model calls and tools.'),('agent/state.py','Defines structured trip state and traveller preferences.'),('agent/schemas.py','Declares tool specifications and argument schemas.'),('agent/prompts.py','Stores conversation instructions and planning prompts.'),('agent/llm_client.py','Creates provider clients and parses one-shot JSON completions.'),('agent/nlu.py','Combines entity extraction and intent results.'),('agent/ner.py','Extracts trip slots using spaCy and regular expressions.'),('agent/intent.py','Trains and applies TF-IDF plus Logistic Regression intent classification.'),('agent/intent_rules.py','Recognises high-precision intents through rules.'),('agent/router.py','Selects NLP or LLM routing using embeddings and state signals.'),('agent/tool_router.py','Checks prerequisites and dispatches deterministic tool calls.'),('agent/tools.py','Implements travel tools and updates trip state.'),('agent/db.py','Saves and reloads SQLite session records.'),('agent/cache.py','Provides disk-cache support for retrieved planning data.')])
page(); h('8.4 Discovery Planning and Route Modules')
table([
('agent/discovery.py','Combines destination sources, filtering and discovery scoring.'),('agent/nearby_search.py','Filters curated destinations by category and radius.'),('agent/destinations_db.py','Stores curated destination names, categories and coordinates.'),('agent/hill_stations.py','Contains hill-station records and specialised nearby search helpers.'),('agent/geo.py','Resolves place names through Ola and Nominatim geocoding.'),('agent/pois.py','Retrieves and parses attractions, stays and food from Overpass.'),('agent/place_profiles.py','Merges approved facts, OSM tags and labelled category defaults.'),('agent/activities.py','Classifies activity duration and effort with validation and caching.'),('agent/scheduler.py','Assigns attractions, orders days and builds validated timelines.'),('agent/route_astar.py','Searches small hotel-to-stops-to-hotel visit sequences.'),('agent/ola_routing.py','Retrieves road distances and route information.'),('agent/ola_elevation.py','Retrieves elevation signals for activity assessment.'),('agent/weather.py','Obtains forecast or historical climate information.'),('agent/trip_conditions.py','Estimates weather, crowd, risk, season and interest suitability.'),('agent/opening_hours.py','Checks supported weekly opening-hours expressions.'),('agent/budget.py','Calculates cost line items from explicit assumptions.'),('agent/transport.py','Estimates intercity transport choices, time and costs.'),('agent/stay.py','Handles accommodation-related planning support.'),('agent/map_view.py','Builds geographic visualisation for the alternative map interface.')])
page(); h('8.5 Backend Frontend and Entry Points')
table([
('server/main.py','Defines health, sessions, chat, ideas, itinerary move and map API endpoints.'),('main.py','Command-line entry point for the agent.'),('app_streamlit.py','Alternative interactive Streamlit application.'),('web/src/main.jsx','Mounts the React application in the browser.'),('web/src/App.jsx','Defines application routing.'),('web/src/api.js','Wraps communication with the backend API.'),('web/src/format.js','Provides display-formatting helpers.'),('web/src/dayColors.js','Defines consistent colours for itinerary days.'),('pages/Landing.jsx','Displays the landing page and trip-session entry interface.'),('pages/Builder.jsx','Coordinates the trip-builder experience.'),('components/ChatPane.jsx','Shows conversation and accepts messages.'),('components/TripPane.jsx','Organises the trip information panel and tabs.'),('components/IdeasTab.jsx','Displays attraction and recommendation ideas.'),('components/ItineraryTab.jsx','Shows daily plans and supports moving itinerary stops.'),('components/MapTab.jsx','Displays selected places and routes on a map.'),('components/StayTab.jsx','Presents accommodation-related information.'),('web/src/index.css','Global styles and frontend styling declarations.'),('web/src/App.css','Application-level styling.')])
p('Page and component paths in this table are relative to web/src. Each interface view consumes backend state rather than duplicating the Python scheduler in the browser.')
page(); h('8.6 Knowledge Maintenance and Test Modules')
table([
('scripts/refresh_official_places.py','Retrieves official-source candidates for human review.'),('data/india_place_profiles.json','Approved attraction facts and provenance.'),('data/official_place_sources.json','Source URLs configured for official-place refresh.'),('official_place_review.json','Generated review queue; used when the refresh script runs.'),('sessions.db','SQLite session store; location depends on the server working directory.'),('test/test_api.py','API-related test coverage.'),('test/test_features.py','Planning feature scenarios.'),('test/test_nlu.py','Entity and intent understanding scenarios.'),('test/test_nlp_tool_router.py','Deterministic tool-dispatch scenarios.'),('test/test_in_between_router.py','Hybrid routing scenarios.'),('test/test_pipeline.py','End-to-end planning pipeline scenarios.'),('test/test_schedule_quality.py','Itinerary feasibility and schedule-quality checks.'),('test/test_multi_select.py','Multiple destination-selection scenarios.'),('test/test_multi_and_vague.py','Multiple and vague destination requests.'),('test/test_pagination.py','Paginated recommendation scenarios.'),('test/test_trip_init_typo.py','Trip initialisation with typo-related inputs.'),('test/test_geonames.py','GeoNames integration checks.'),('test/test_llm_budget.py','Language-model and budget interaction checks.'),('test/test_nvidia.py','NVIDIA provider integration checks.'),('probe_overpass_status.py','Standalone Overpass-service diagnostic script.')])

screens=[
('Landing page and new trip button','Landing page of the AI Travel Agent',
'The landing page introduces the travel planner and provides the starting point for a new trip. The traveller opens the planning interface through the new trip button. This screen connects the project overview with the conversational workflow and gives users a clear way to begin entering their travel requirements.'),
('Saved trip cards on the landing page','Saved trip sessions',
'Saved trip cards show the destination, origin, dates and current planning stage for earlier sessions. Selecting a card reopens its conversation and trip details. This output illustrates how SQLite persistence supports continued planning, allowing the traveller to return to a trip without entering the same details again.'),
('Chat showing entered trip details','Conversational collection of trip preferences',
'The chat view shows a traveller entering dates, destination, budget and group size in natural language. The agent extracts these details into structured trip state and asks for missing information. This output illustrates how entity extraction supports a conversation that gathers the requirements needed for a travel plan.'),
('Chat recommending nearby destinations','Nearby destination recommendations',
'The conversation presents destination candidates for a category request such as nearby hill stations. Geographic distance and curated categories help identify relevant choices. The traveller can select a suggestion to continue planning, illustrating how the system handles a trip request before a destination is fixed.'),
('Ideas tab with attraction recommendations','Attraction recommendations',
'The Ideas tab lists attractions retrieved for the selected destination and provides information for comparing places. The traveller can review the available options before the itinerary is built. This output connects geographic data and place profiles with the set of candidate activities considered by the scheduler.'),
('Chat response showing weather information','Weather information for the trip dates',
'The weather response presents conditions for the selected destination and travel dates. It identifies whether the information is a forecast or a historical climate estimate. This output helps the traveller review outdoor activity suitability and understand the weather assumptions that influence the suggested itinerary.'),
('Chat response comparing transport options','Transport options and travel estimates',
'The transport response compares supported travel modes between the origin and selected destination. Estimated journey time and cost help the traveller assess practical options. This output shows how transport planning uses distance and mode assumptions to support a decision without claiming confirmed fares or bookings.'),
('Generated day-wise itinerary with timings','Personalised day-wise travel itinerary',
'The itinerary view groups selected attractions into trip days and displays estimated arrival and departure times. Visit duration, travel time and daily effort limits shape the schedule. This output shows how the planner converts recommended places into an organised sequence of activities that the traveller can review.'),
('Interactive map displaying itinerary places and route','Map visualisation of travel locations',
'The map displays itinerary stops at their geographic coordinates and uses day colours to distinguish groups of visits. Route information helps the traveller compare locations with the daily plan. This output makes the distance and sequence relationships visible, supporting a review of how the selected places fit together.'),
('Stay tab showing accommodation recommendations','Accommodation recommendations',
'The Stay tab presents accommodation information for the selected destination. The traveller can compare the details alongside planned attractions and dates. This output shows how stay recommendations form part of the trip workflow, helping the traveller review available information before arranging a confirmed booking.'),
('Costs tab with total estimate and cost breakdown','Estimated trip cost breakdown',
'The Costs tab displays the total trip estimate and its accommodation, food, local travel and activity components. The values use stated assumptions when verified prices are unavailable. This output helps the traveller compare spending categories and assess the proposed plan against the available budget before making bookings.'),
('Itinerary after moving a stop between days','Manual itinerary editing',
'The itinerary view shows a stop after the traveller moves it to another day or changes its position. The backend checks the requested edit and recalculates the affected timeline. This output illustrates how the generated plan can be adjusted through the interface while retaining the scheduling checks used by the application.'),
('Itinerary details showing intensity and route source','Planning constraints and source information',
'The itinerary shows daily intensity, effort and route sources alongside scheduled activities. These fields explain the limits and data assumptions used by the planner. This output helps the traveller distinguish road data from fallback estimates and review the practical reasons behind the arrangement of each trip day.'),
('Unscheduled attractions and omission reasons','Unscheduled attractions and reasons',
'The unscheduled section lists attractions omitted from the itinerary and explains each omission. Reasons include distance restrictions or limits on the available activity pool. This output helps the traveller understand why a suggested place is absent and review which recommendations could not fit into the generated plan.'),
('Trace tab showing planning tool activity','Planning execution trace',
'The Trace tab presents tool activity recorded during the trip conversation. It explains which planning operations were invoked as the user supplied details and requested outputs. This screen provides a view of the agent workflow, helping the reviewer follow how the application produced its recommendations and itinerary.'),
('Chat selecting more than one destination','Multiple destination selection',
'The conversation shows the traveller choosing more than one destination from the suggested candidates. The agent records the selection and continues the trip dialogue using the updated state. This output illustrates how a single planning session can collect several destinations while preserving the original travel requirements.'),
('Food ideas panel listing restaurants and cuisines','Food and restaurant recommendations',
'The food ideas panel lists dining places retrieved near the chosen destination and shows cuisine details when available. These suggestions help the traveller review meal options alongside sightseeing plans. This output illustrates how restaurant information complements attraction recommendations within the same trip workflow.'),
('Ideas tab showing included and excluded attractions','Attraction inclusion and exclusion',
'The Ideas tab distinguishes included attractions from places excluded by the traveller. Selecting an entry changes its status and requests an updated itinerary. This output shows how direct user preferences affect the planning pool, allowing the traveller to remove unwanted stops without starting a new trip conversation.'),
('Chat requesting a relaxed pace and the updated itinerary','Changing the itinerary pace',
'The conversation shows a request to make the trip more relaxed and the resulting itinerary update. The planner applies the selected pace to its daily activity limits. This output illustrates how a short natural-language instruction can revise the workload of a trip while keeping the existing destination and travel dates.'),
('Daily itinerary showing meal stops and timings','Meal placement in the daily timeline',
'The daily timeline includes meal stops alongside attraction visits and travel intervals. Their placement helps the traveller review how sightseeing and meals fit within the available day. This output illustrates the scheduler’s timeline construction, including the time reserved for activities beyond the main attraction visits.'),
('Map with an attraction marker popup open','Attraction details on the map',
'An open map marker popup identifies a selected attraction within its geographic setting. The surrounding markers help the traveller relate this place to other stops on the route. This output illustrates how the map supports closer inspection of individual locations alongside the broader view of the planned daily journey.'),
('Attraction card showing entry fees opening hours and source','Place facts and source provenance',
'An attraction card shows available entry fees, opening hours and source information next to the planned visit. These details help the traveller review practical conditions before travelling. This output illustrates how the interface presents sourced facts alongside estimates, with missing information requiring further verification.'),
('Trace entry showing extracted entities and intent','Extracted trip slots and intent',
'The NLU trace entry displays the recognised intent and trip details extracted from a user message. It connects the wording in the conversation with structured values used by the planner. This output helps the reviewer inspect entity extraction and intent classification without treating a displayed prediction as measured accuracy.'),
('Trace entry showing the NLP or LLM routing decision','NLP and LLM routing decision',
'A routing trace entry records whether the message was directed to the NLP tool engine or the LLM path. The available decision details help explain how the hybrid router selected its next action. This output illustrates the link between language understanding, dialogue state and the planning operation used to answer the request.'),
('Itinerary showing a validation message for a rejected move','Validation of an itinerary edit',
'The itinerary displays a validation message when a requested stop movement cannot be accepted. The traveller can review the response and choose a different adjustment. This output illustrates how editing checks protect the plan from unsuitable changes, rather than accepting every movement without checking the affected schedule.')]
chapter(9,topics[8])
for i,(label,title,description) in enumerate(screens,1):
    if i>1: page()
    h(f'9.{i} {title}')
    gap(f'[Insert screenshot of {label.lower()} here]',13)
    caption(f'Figure 9.{i}  {title}')
    a=d.add_paragraph(); a.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    a.paragraph_format.line_spacing=1.5
    spacing(a.add_run(description))

chapter(10,topics[9]); h('10.1 Conclusion')
p('The AI Travel Agent brings conversational input, geographic data and itinerary optimisation into one application. NLP modules extract travel requirements and recognise intents, while the hybrid router selects direct tool execution or LLM reasoning according to the message and dialogue state. The planning tools retrieve destinations and attractions, estimate activity properties and construct day-wise schedules through CP-SAT.')
p('Route ordering and map views make the resulting plan easier to inspect. SQLite persistence lets the user continue an existing trip, and provenance labels help distinguish approved facts, external tags and estimates. The project demonstrates how language understanding can connect to deterministic algorithms that enforce explicit constraints rather than relying on a generated paragraph alone.')
h('10.2 Current Limitations')
p('The planner depends on the completeness of source data and the availability of external services. Curated attraction coverage is limited, complex opening-hour expressions are not fully verified, and activity, crowd and budget values may be estimates. Long-range weather uses historical climate data. Route quality depends on the available distance matrix, and the scheduler’s feasibility model does not guarantee booking availability or total budget compliance.')
h('10.3 Future Enhancement')
p('Future work can add multilingual entity extraction and a larger labelled intent dataset evaluated on a separate test set. More manually verified attraction profiles would improve duration, fees, accessibility and opening-hours coverage. Live hotel and ticket integrations could replace assumption-based prices with available quotations and support explicit budget constraints.')
p('Further improvements include traffic-aware travel times, feedback-based recommendations, collaborative group planning and mobile access. Evaluation should measure entity extraction, intent classification, route cost, constraint violations and completed user journeys. These measurements would provide evidence for improvements beyond the current implementation description.')

page(); d.add_paragraph('REFERENCES AND SOURCE MATERIALS','Heading 1')
p('The implementation descriptions and numerical assumptions in this report are based on the project source files listed below. Package names identify the software used; no performance measurements are inferred from the existence of a dependency or test file.')
refs=[
'Project README.md and requirements.txt: application overview, dependency declarations and data-source conventions.',
'web/package.json and web/src: React interface, mapping packages and frontend component structure.',
'server/main.py: session, chat, recommendation, map and itinerary-edit API endpoints.',
'agent/ner.py, agent/nlu.py, agent/intent.py and agent/intent_rules.py: entity rules, intent training examples, TF-IDF and confidence handling.',
'agent/router.py and agent/tool_router.py: semantic routing, TF-IDF fallback and deterministic dispatch.',
'agent/orchestrator.py, agent/llm_client.py, agent/schemas.py and agent/prompts.py: dialogue flow and language-model integration.',
'agent/scheduler.py and agent/route_astar.py: attraction assignment, objective terms, route sequencing and timeline checks.',
'agent/budget.py and agent/transport.py: documented cost assumptions and transport estimates.',
'agent/pois.py, agent/place_profiles.py, agent/weather.py and agent/trip_conditions.py: external records, provenance and suitability estimates.',
'agent/db.py, agent/state.py and scripts/refresh_official_places.py: persistence, state models and reviewed knowledge maintenance.'
]
for i,r in enumerate(refs,1): p(f'[{i}] {r}')

footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run(); r.font.name='Times New Roman'; r.font.size=Pt(10)
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); r._r.addnext(f)
d.core_properties.title='AI Travel Agent College Project Report'
d.core_properties.subject='Artificial Intelligence and Natural Language Processing'
d.core_properties.author=''
d.core_properties.keywords='travel planning, NLP, CP-SAT, A star, React, FastAPI'
d.save(OUT)
print(OUT)
print('Paragraphs:',len(d.paragraphs),'Tables:',len(d.tables))
